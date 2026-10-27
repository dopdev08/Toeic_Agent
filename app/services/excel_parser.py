import io
from typing import Any, BinaryIO
import openpyxl
from pydantic import ValidationError
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models.question import Question
from app.schemas.question import (
    AnswerKeyConfirmationItem,
    ImportRowError,
    PreviewSummary,
    QuestionChangeComparison,
    QuestionFieldDiff,
    QuestionImportPreviewResponse,
    QuestionImportResponse,
    QuestionItem,
    QuestionPreviewItem,
)

REQUIRED_COLUMNS = [
    "test_id",
    "part",
    "question_no",
    "question_text",
    "option_a",
    "option_b",
    "option_c",
    "option_d",
    "correct_answer",
]


class ExcelImportError(Exception):
    """Raised when Excel parsing or business validation fails."""
    pass


def validate_file_extension(filename: str | None) -> None:
    if not filename or not filename.lower().endswith(".xlsx"):
        raise ExcelImportError("File must have a .xlsx extension")


def _is_row_empty(row_values: list[object]) -> bool:
    return all(
        val is None or (isinstance(val, str) and not val.strip())
        for val in row_values
    )


def _format_row_error(row_idx: int, error: ValidationError) -> str:
    err = error.errors()[0]
    field = str(err["loc"][0])

    if field == "correct_answer":
        return f"Row {row_idx}: correct_answer must be one of A, B, C, D"
    if field == "part":
        return f"Row {row_idx}: part must be an integer from 1 to 7"
    if field == "question_no":
        return f"Row {row_idx}: question_no must be a positive integer"
    return f"Row {row_idx}: {field} {err['msg']}"


def parse_and_validate_questions(file_content: bytes | BinaryIO) -> QuestionImportResponse:
    if isinstance(file_content, bytes):
        stream = io.BytesIO(file_content)
    else:
        stream = file_content

    try:
        workbook = openpyxl.load_workbook(stream, data_only=True)
    except Exception as exc:
        raise ExcelImportError("Invalid or corrupted Excel file") from exc

    sheet = workbook.active
    if sheet is None:
        raise ExcelImportError("Excel file has no active sheet")

    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        raise ExcelImportError("Excel file is empty")

    header_row = rows[0]
    col_mapping: dict[str, int] = {}
    for idx, cell in enumerate(header_row):
        if cell is not None and str(cell).strip():
            normalized_col = str(cell).strip().lower()
            col_mapping[normalized_col] = idx

    missing_cols = [col for col in REQUIRED_COLUMNS if col not in col_mapping]
    if missing_cols:
        raise ExcelImportError(f"Missing required columns: {', '.join(missing_cols)}")

    questions: list[QuestionItem] = []
    seen_question_numbers: set[tuple[str, int]] = set()
    first_test_id: str | None = None

    for row_offset, row_cells in enumerate(rows[1:], start=2):
        if _is_row_empty(list(row_cells)):
            continue

        row_data: dict[str, object] = {}
        for col_name in REQUIRED_COLUMNS:
            col_idx = col_mapping[col_name]
            val = row_cells[col_idx] if col_idx < len(row_cells) else None

            if val is None or (isinstance(val, str) and not val.strip()):
                raise ExcelImportError(f"Row {row_offset}: {col_name} cannot be empty")

            row_data[col_name] = val

        if "explanation" in col_mapping:
            exp_idx = col_mapping["explanation"]
            exp_val = row_cells[exp_idx] if exp_idx < len(row_cells) else None
            if exp_val is not None and str(exp_val).strip():
                row_data["explanation"] = str(exp_val).strip()

        try:
            item = QuestionItem.model_validate(row_data)
        except ValidationError as val_err:
            raise ExcelImportError(_format_row_error(row_offset, val_err)) from val_err

        if first_test_id is None:
            first_test_id = item.test_id
        elif item.test_id != first_test_id:
            raise ExcelImportError(
                f"Row {row_offset}: test_id '{item.test_id}' does not match expected '{first_test_id}'"
            )

        key = (item.test_id, item.question_no)
        if key in seen_question_numbers:
            raise ExcelImportError(
                f"Row {row_offset}: duplicate question_no {item.question_no} for test_id '{item.test_id}'"
            )
        seen_question_numbers.add(key)
        questions.append(item)

    if not questions:
        raise ExcelImportError("Excel file contains no question data")

    return QuestionImportResponse(
        test_id=first_test_id,
        total_questions=len(questions),
        imported_count=0,
        skipped_count=0,
        questions=questions,
    )


def generate_import_preview(
    file_content: bytes | BinaryIO,
    filename: str,
    db: Session,
) -> QuestionImportPreviewResponse:
    """Parse Excel workbook in read-only mode and compute a preview diff against DB.

    - Never writes, flushes, or commits to the database.
    - Captures row-level validation errors rather than aborting at the first failure.
    - Detects within-file duplicates and gaps in question numbering.
    - Compares each valid question to existing database records on (test_id, part, question_no).
    - Classifies actions as INSERT, UNCHANGED, or UPDATE (conflict).
    - Flags answer key changes requiring authoritative verification.
    """
    if isinstance(file_content, bytes):
        stream = io.BytesIO(file_content)
    else:
        stream = file_content

    validate_file_extension(filename)

    try:
        workbook = openpyxl.load_workbook(stream, data_only=True)
    except Exception as exc:
        raise ExcelImportError("Invalid or corrupted Excel file") from exc

    sheet = None
    if "questions" in workbook.sheetnames:
        sheet = workbook["questions"]
    else:
        sheet = workbook.active

    if sheet is None:
        raise ExcelImportError("Excel file has no active sheet")

    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        raise ExcelImportError("Excel file is empty")

    header_row = rows[0]
    col_mapping: dict[str, int] = {}
    for idx, cell in enumerate(header_row):
        if cell is not None and str(cell).strip():
            normalized_col = str(cell).strip().lower()
            col_mapping[normalized_col] = idx

    missing_cols = [col for col in REQUIRED_COLUMNS if col not in col_mapping]
    validation_errors: list[ImportRowError] = []

    if missing_cols:
        validation_errors.append(
            ImportRowError(
                row_no=1,
                field="headers",
                error=f"Missing required columns: {', '.join(missing_cols)}",
            )
        )
        return QuestionImportPreviewResponse(
            filename=filename,
            summary=PreviewSummary(
                total_rows=max(0, len(rows) - 1),
                valid_questions=0,
                invalid_rows=max(0, len(rows) - 1),
                to_insert=0,
                unchanged=0,
                to_update=0,
                conflicts=0,
                duplicates_in_file=0,
            ),
            validation_errors=validation_errors,
        )

    total_data_rows = 0
    valid_items: list[tuple[int, QuestionItem]] = []
    seen_in_file: dict[tuple[str, int, int], int] = {}
    duplicates_in_file: list[dict[str, Any]] = []
    preview_items: list[QuestionPreviewItem] = []

    for row_offset, row_cells in enumerate(rows[1:], start=2):
        if _is_row_empty(list(row_cells)):
            continue

        total_data_rows += 1
        row_errors: list[ImportRowError] = []
        row_dict: dict[str, Any] = {}

        raw_qno = None
        for col_name in REQUIRED_COLUMNS:
            idx = col_mapping[col_name]
            val = row_cells[idx] if idx < len(row_cells) else None

            if val is None or (isinstance(val, str) and not val.strip()):
                row_errors.append(
                    ImportRowError(
                        row_no=row_offset,
                        field=col_name,
                        error=f"{col_name} cannot be empty",
                    )
                )
            else:
                row_dict[col_name] = val
                if col_name == "question_no":
                    try:
                        raw_qno = int(val)
                    except (ValueError, TypeError):
                        pass

        if "explanation" in col_mapping:
            exp_idx = col_mapping["explanation"]
            exp_val = row_cells[exp_idx] if exp_idx < len(row_cells) else None
            if exp_val is not None and str(exp_val).strip():
                row_dict["explanation"] = str(exp_val).strip()

        for err in row_errors:
            err.question_no = raw_qno

        item: QuestionItem | None = None
        if not row_errors:
            try:
                item = QuestionItem.model_validate(row_dict)
            except ValidationError as val_err:
                for err in val_err.errors():
                    field_name = str(err["loc"][0])
                    msg = err["msg"]
                    if field_name == "correct_answer":
                        msg = "must be one of A, B, C, D"
                    raw_val = str(row_dict.get(field_name, ""))
                    row_errors.append(
                        ImportRowError(
                            row_no=row_offset,
                            question_no=raw_qno,
                            field=field_name,
                            error=f"{field_name} {msg}",
                            raw_value=raw_val if raw_val else None,
                        )
                    )

        if row_errors:
            validation_errors.extend(row_errors)
            preview_items.append(
                QuestionPreviewItem(
                    row_no=row_offset,
                    question_no=raw_qno or 0,
                    test_id=str(row_dict.get("test_id", "")),
                    part=int(row_dict.get("part", 5)) if str(row_dict.get("part", "")).isdigit() else 5,
                    question_text=str(row_dict.get("question_text", "")),
                    option_a=str(row_dict.get("option_a", "")),
                    option_b=str(row_dict.get("option_b", "")),
                    option_c=str(row_dict.get("option_c", "")),
                    option_d=str(row_dict.get("option_d", "")),
                    correct_answer=str(row_dict.get("correct_answer", "")),
                    explanation=str(row_dict.get("explanation", "")) if row_dict.get("explanation") else None,
                    action="INVALID",
                )
            )
            continue

        key = (item.test_id, item.part, item.question_no)
        if key in seen_in_file:
            first_seen_row = seen_in_file[key]
            dup_info = {
                "test_id": item.test_id,
                "part": item.part,
                "question_no": item.question_no,
                "first_row": first_seen_row,
                "duplicate_row": row_offset,
            }
            duplicates_in_file.append(dup_info)
            validation_errors.append(
                ImportRowError(
                    row_no=row_offset,
                    question_no=item.question_no,
                    field="question_no",
                    error=f"Duplicate question_no {item.question_no} for test_id '{item.test_id}' in workbook (first seen at row {first_seen_row})",
                )
            )
            preview_items.append(
                QuestionPreviewItem(
                    row_no=row_offset,
                    question_no=item.question_no,
                    test_id=item.test_id,
                    part=item.part,
                    question_text=item.question_text,
                    option_a=item.option_a,
                    option_b=item.option_b,
                    option_c=item.option_c,
                    option_d=item.option_d,
                    correct_answer=item.correct_answer,
                    explanation=item.explanation,
                    action="INVALID",
                )
            )
            continue

        seen_in_file[key] = row_offset
        valid_items.append((row_offset, item))

    detected_test_ids = {item.test_id for _, item in valid_items}
    detected_parts = {item.part for _, item in valid_items}
    detected_test_id = valid_items[0][1].test_id if len(detected_test_ids) == 1 else None
    detected_part = valid_items[0][1].part if len(detected_parts) == 1 else None

    q_numbers = [item.question_no for _, item in valid_items]
    min_qno = min(q_numbers) if q_numbers else None
    max_qno = max(q_numbers) if q_numbers else None
    missing_question_numbers: list[int] = []
    if min_qno is not None and max_qno is not None:
        expected_range = set(range(min_qno, max_qno + 1))
        missing_question_numbers = sorted(list(expected_range - set(q_numbers)))

    existing_db_records = (
        db.scalars(select(Question).where(Question.test_id.in_(detected_test_ids))).all()
        if detected_test_ids
        else []
    )
    existing_map: dict[tuple[str, int, int], Question] = {
        (q.test_id, q.part, q.question_no): q for q in existing_db_records
    }

    to_insert_count = 0
    unchanged_count = 0
    to_update_count = 0
    conflicts_count = 0
    changes: list[QuestionChangeComparison] = []
    requires_confirmation: list[AnswerKeyConfirmationItem] = []

    for row_offset, item in valid_items:
        key = (item.test_id, item.part, item.question_no)
        existing = existing_map.get(key)

        if existing is None:
            action = "INSERT"
            to_insert_count += 1
            diffs: list[QuestionFieldDiff] = []
        else:
            diffs = []
            if existing.question_text != item.question_text:
                diffs.append(
                    QuestionFieldDiff(
                        field="question_text",
                        old_value=existing.question_text,
                        new_value=item.question_text,
                    )
                )
            if existing.option_a != item.option_a:
                diffs.append(
                    QuestionFieldDiff(
                        field="option_a",
                        old_value=existing.option_a,
                        new_value=item.option_a,
                    )
                )
            if existing.option_b != item.option_b:
                diffs.append(
                    QuestionFieldDiff(
                        field="option_b",
                        old_value=existing.option_b,
                        new_value=item.option_b,
                    )
                )
            if existing.option_c != item.option_c:
                diffs.append(
                    QuestionFieldDiff(
                        field="option_c",
                        old_value=existing.option_c,
                        new_value=item.option_c,
                    )
                )
            if existing.option_d != item.option_d:
                diffs.append(
                    QuestionFieldDiff(
                        field="option_d",
                        old_value=existing.option_d,
                        new_value=item.option_d,
                    )
                )
            if existing.correct_answer != item.correct_answer:
                diffs.append(
                    QuestionFieldDiff(
                        field="correct_answer",
                        old_value=existing.correct_answer,
                        new_value=item.correct_answer,
                    )
                )
            if (existing.explanation or "") != (item.explanation or ""):
                diffs.append(
                    QuestionFieldDiff(
                        field="explanation",
                        old_value=existing.explanation,
                        new_value=item.explanation,
                    )
                )

            if not diffs:
                action = "UNCHANGED"
                unchanged_count += 1
            else:
                action = "UPDATE"
                to_update_count += 1
                conflicts_count += 1
                is_answer_change = any(d.field == "correct_answer" for d in diffs)
                changes.append(
                    QuestionChangeComparison(
                        test_id=item.test_id,
                        part=item.part,
                        question_no=item.question_no,
                        existing_id=existing.id,
                        diffs=diffs,
                        old_correct_answer=existing.correct_answer,
                        new_correct_answer=item.correct_answer,
                        is_answer_change=is_answer_change,
                    )
                )
                if is_answer_change:
                    requires_confirmation.append(
                        AnswerKeyConfirmationItem(
                            question_no=item.question_no,
                            test_id=item.test_id,
                            part=item.part,
                            old_answer=existing.correct_answer,
                            new_answer=item.correct_answer,
                            question_text=item.question_text,
                            reason=f"Proposed correct answer '{item.correct_answer}' differs from stored answer '{existing.correct_answer}' and requires confirmation against authoritative key.",
                        )
                    )

        preview_items.append(
            QuestionPreviewItem(
                row_no=row_offset,
                question_no=item.question_no,
                test_id=item.test_id,
                part=item.part,
                question_text=item.question_text,
                option_a=item.option_a,
                option_b=item.option_b,
                option_c=item.option_c,
                option_d=item.option_d,
                correct_answer=item.correct_answer,
                explanation=item.explanation,
                action=action,
                diffs=diffs,
            )
        )

    preview_items.sort(key=lambda p: p.row_no)

    summary = PreviewSummary(
        total_rows=total_data_rows,
        valid_questions=len(valid_items),
        invalid_rows=total_data_rows - len(valid_items),
        to_insert=to_insert_count,
        unchanged=unchanged_count,
        to_update=to_update_count,
        conflicts=conflicts_count,
        duplicates_in_file=len(duplicates_in_file),
    )

    return QuestionImportPreviewResponse(
        filename=filename,
        test_id=detected_test_id,
        part=detected_part,
        min_question_no=min_qno,
        max_question_no=max_qno,
        missing_question_numbers=missing_question_numbers,
        summary=summary,
        validation_errors=validation_errors,
        duplicates_in_file=duplicates_in_file,
        changes=changes,
        requires_confirmation=requires_confirmation,
        questions=preview_items,
    )

